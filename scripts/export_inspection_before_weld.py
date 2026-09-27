# 용접 전에 검사가 기록된 조인트(용접일 없음 / 검사일이 용접일보다 빠름)를 엑셀로 출력하는 스크립트
import os
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from supabase import create_client, ClientOptions

INSP = [("VT", "vt_date", "vt_result"), ("RT", "rt_date", "rt_result"), ("RT 2nd", "rt_2_date", "rt_2_result"),
        ("PT", "pt_date", "pt_result"), ("MT", "mt_date", "mt_result"), ("PWHT", "pwht_date", "pwht_result")]
BASE = ["system", "unit", "area", "sub_area", "line_no", "iso_drawing", "rev", "joint_no", "size_inch", "mat",
        "package", "welder", "date_completed", "inspection", "pwht"]
COLS = ",".join(["id"] + BASE + [c for _, d, r in INSP for c in (d, r)])
RED = PatternFill("solid", fgColor="FFC7CE")
HEAD = PatternFill("solid", fgColor="1E3A8A")


def load_env():
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    with open(os.path.join(base_dir, ".env"), "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                k, v = line.split("=", 1)
                os.environ[k.strip()] = v.strip()
    return base_dir


def fetch_all(sb):
    total = sb.table("joint_master").select("id", count="exact").limit(1).execute().count
    rows, off = [], 0
    while True:
        b = sb.table("joint_master").select(COLS).order("id").range(off, off + 9999).execute().data or []
        rows.extend(b)
        if len(b) < 10000:
            break
        off += 10000
    if len(rows) != total:
        raise RuntimeError(f"joint_master {len(rows)}건을 읽었지만 테이블은 {total}건입니다. 다시 실행하세요.")
    return rows


def has(v):
    return v not in (None, "")


def write_sheet(ws, rows, flag_fn, note_fn):
    head = [c.upper() for c in BASE] + [h for n, _, _ in INSP for h in (f"{n} DATE", f"{n} RESULT")] + ["NOTE"]
    ws.append(head)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
    for r in rows:
        ws.append([r.get(c) for c in BASE] + [r.get(c) for _, d, rs in INSP for c in (d, rs)] + [note_fn(r)])
        for i, (_, d, rs) in enumerate(INSP):
            if flag_fn(r, d, rs):
                col = len(BASE) + 1 + i * 2
                ws.cell(ws.max_row, col).fill = RED
                ws.cell(ws.max_row, col + 1).fill = RED
    for i, h in enumerate(head, 1):
        width = max([len(str(h))] + [len(str(ws.cell(j, i).value or "")) for j in range(2, ws.max_row + 1)])
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 45)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    base_dir = load_env()
    sb = create_client(os.environ["SUPABASE_URL"], os.environ["SUPABASE_KEY"], options=ClientOptions(schema="construction"))
    rows = fetch_all(sb)

    # 1) 용접일이 없는데 검사 날짜/결과가 입력된 조인트
    not_welded = [r for r in rows if not r.get("date_completed") and any(has(r.get(d)) or has(r.get(rs)) for _, d, rs in INSP)]
    # 2) 검사일이 용접일보다 빠른 조인트(앱 QA_CHECKS["date_error"]와 같은 조건)
    weld = lambda r: str(r["date_completed"])[:10]
    early = [r for r in rows if r.get("date_completed") and any(has(r.get(d)) and str(r[d])[:10] < weld(r) for _, d, _ in INSP)]
    key = lambda r: (r.get("iso_drawing") or "", str(r.get("joint_no") or ""))
    not_welded.sort(key=key)
    early.sort(key=key)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inspection before Weld"
    write_sheet(ws, early,
                lambda r, d, rs: has(r.get(d)) and str(r[d])[:10] < weld(r),
                lambda r: ", ".join(f"{n} {str(r[d])[:10]} < Weld {weld(r)}" for n, d, _ in INSP if has(r.get(d)) and str(r[d])[:10] < weld(r)))
    write_sheet(wb.create_sheet("Not Welded with Inspection"), not_welded,
                lambda r, d, rs: has(r.get(d)) or has(r.get(rs)),
                lambda r: "Weld date empty, inspection recorded: " + ", ".join(n for n, d, rs in INSP if has(r.get(d)) or has(r.get(rs))))

    out = os.path.join(base_dir, "Reports", f"Inspection_before_Weld_{datetime.now():%Y%m%d}.xlsx")
    wb.save(out)
    print(f"전체 {len(rows)}건 중 검사일<용접일 {len(early)}건, 용접일 없이 검사 기록 {len(not_welded)}건 → {out}")


if __name__ == "__main__":
    main()
